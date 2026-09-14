#!/usr/bin/env python3
"""
Build Pilot Stock Snapshot from stock(1).xlsx
Extracts P/N, On Hand (F1 & F2), categories, specs, and product colors.
Output: assets/js/pilot-stock-snapshot.js
"""

import os
import sys
import json
import re
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

ROOT_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
EXCEL_PATH = os.path.join(ROOT_DIR, 'stock(1).xlsx')
OUTPUT_JS = os.path.join(ROOT_DIR, 'assets', 'js', 'pilot-stock-snapshot.js')

KNOWN_COLORS = [
    "Titanium Silverblue",
    "Titanium Black",
    "Titanium Gray",
    "Titanium White",
    "Cobalt Violet",
    "Light Violet",
    "Light Blue",
    "Blue Violet",
    "Pink Gold",
    "Pistachio",
    "Graphite",
    "Blueberry",
    "Jetblack",
    "Icyblue",
    "Silver",
    "White",
    "Black",
    "Gray",
    "Grey",
    "Violet",
    "Lavender",
    "Cream",
    "Pink",
    "Mint",
    "Navy",
    "Coral Red",
    "Dark Green",
    "Dark Blue",
    "Sky Blue"
]

COLOR_CANONICAL_NAMES = {
    "navy": "Navy",
    "jetblack": "Jet Black",
    "icyblue": "Icy Blue",
    "lightviolet": "Light Violet",
    "light violet": "Light Violet",
    "titaniumblack": "Titanium Black",
    "titanium black": "Titanium Black",
    "titaniumsilverblue": "Titanium Silverblue",
    "titanium silverblue": "Titanium Silverblue",
    "graphite": "Graphite",
    "pistachio": "Pistachio",
    "blueberry": "Blueberry"
}

def normalize_color_name(value):
    raw = str(value or "").strip()
    if not raw:
        return ""
    key = re.sub(r'\s+', ' ', raw.lower())
    compact = key.replace(' ', '')
    if key in COLOR_CANONICAL_NAMES:
        return COLOR_CANONICAL_NAMES[key]
    if compact in COLOR_CANONICAL_NAMES:
        return COLOR_CANONICAL_NAMES[compact]
    return raw.title()

def extract_color_from_description(description):
    text = str(description or "").strip()
    if not text:
        return ""
    m = re.search(r'\s*-\s*([^-]+)\s*$', text)
    if not m:
        return ""
    candidate = m.group(1).strip()
    if not candidate or re.match(r'^\d', candidate) or re.match(r'^(5G|4G|LTE|WI-?FI)$', candidate, re.IGNORECASE):
        return ""
    return candidate

def extract_known_color(description):
    text = str(description or "").strip().lower()
    for color in sorted(KNOWN_COLORS, key=lambda c: -len(c)):
        if text.endswith(color.lower()):
            return color
    return ""

def resolve_product_color(description):
    extracted = extract_color_from_description(description) or extract_known_color(description) or ""
    return normalize_color_name(extracted)

def parse_sheet(ws, sheet_name):
    rows = []
    for r in range(4, ws.max_row + 1):
        cat1 = ws.cell(row=r, column=1).value
        cat2 = ws.cell(row=r, column=2).value
        cat3 = ws.cell(row=r, column=3).value
        brand = ws.cell(row=r, column=4).value
        srp = ws.cell(row=r, column=5).value
        pn = ws.cell(row=r, column=6).value
        sku = ws.cell(row=r, column=7).value
        apple = ws.cell(row=r, column=8).value
        desc = ws.cell(row=r, column=9).value
        on_hand = ws.cell(row=r, column=10).value
        
        raw_pn = str(pn or '').strip()
        if not raw_pn or raw_pn.upper() == 'TOTAL' or str(cat1 or '').upper() == 'TOTAL':
            continue
            
        try:
            qty = int(on_hand or 0)
        except:
            qty = 0
            
        rows.append({
            'pn': raw_pn.upper(),
            'cat1': str(cat1 or '').strip().upper(),
            'cat2': str(cat2 or '').strip().upper(),
            'cat3': str(cat3 or '').strip().upper(),
            'brand': str(brand or '').strip().upper(),
            'desc': str(desc or '').strip().upper(),
            'raw_desc': str(desc or '').strip(),
            'srp': float(srp or 0),
            'onHand': qty
        })
    return rows

def classify_user_rule(item):
    c1 = item['cat1']
    c2 = item['cat2']
    c3 = item['cat3']
    brand = item['brand']
    desc = item['desc']
    pn = item['pn']
    
    # 1. Galaxy Buds (MUST be before Smartphone)
    if (c1 == 'AUDIO' and c2 == 'HEADPHONE' and c3 == 'TRUE WIRELESS' and 'SAMSUNG' in brand) or \
       ('SAMSUNG' in brand and (pn.startswith(('SM-R4', 'SM-R5', 'SM-R6')) or 'BUDS' in desc)):
        return 'Buds'
        
    # 2. Smartphone
    if c1 in ('SMART PHONES', 'SMARTPHONES', 'SMART PHONE', 'SMART_PHONES', 'SMART_PHONE'):
        return 'SmartPhone'
        
    # 3. Tablet
    if c1 in ('COMPUTER AND TABLET', 'COMPUTER_AND_TABLET', 'TABLET', 'TABLETS', 'TAB'):
        return 'Tablet'
        
    # 4. Smart Watch
    if c1 in ('SMART WATCH', 'SMART_WATCH', 'SMARTWATCH', 'WATCH'):
        return 'Watch'
        
    # 5. Accessories
    if c1 in ('MOBILE AND COMPUTER ACCESSORY', 'MOBILE_AND_COMPUTER_ACCESSORY', 'ACCESSORY', 'ACCESSORIES', 'ADAPTER'):
        return 'Accessory'
        
    # 6. Premium
    if 'PREMIUM' in c1 or 'PREMIUM' in c2 or 'FREE GIFT' in c2 or 'PREMIUM' in desc or 'FREE GIFT' in desc or 'GIFT' in c1:
        return 'Premium'
        
    # 7. SIM
    if 'SERVICE, INSURANCE AND WARRANTY' in c1 or 'SERVICE,_INSURANCE_AND_WARRANTY' in c1 or 'SIM' in c1 or 'SIM' in c2 or 'CARRIER MOBILE PACKAGE' in c2:
        return 'SIM'
        
    return 'Other'

def main():
    if not os.path.exists(EXCEL_PATH):
        print(f"Error: {EXCEL_PATH} not found")
        sys.exit(1)
        
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    ws1 = wb['Sheet1']
    ws2 = wb['Sheet2']
    
    s1_rows = parse_sheet(ws1, 'Sheet1')
    s2_rows = parse_sheet(ws2, 'Sheet2')
    
    s1_map = {r['pn']: r for r in s1_rows}
    s2_map = {r['pn']: r for r in s2_rows}
    all_pns = sorted(list(set(list(s1_map.keys()) + list(s2_map.keys()))))
    
    merged_items = []
    f1_total = 0
    f2_total = 0
    
    for pn in all_pns:
        r1 = s1_map.get(pn)
        r2 = s2_map.get(pn)
        f1 = r1['onHand'] if r1 else 0
        f2 = r2['onHand'] if r2 else 0
        ref = r1 or r2
        cat = classify_user_rule(ref)
        color = resolve_product_color(ref['raw_desc'])
        
        f1_total += f1
        f2_total += f2
        
        merged_items.append({
            'pn': pn,
            'model': ref['raw_desc'] or pn,
            'description': ref['raw_desc'] or pn,
            'color': color,
            'category': cat,
            'canonicalCategory': cat,
            'category1': ref['cat1'],
            'category2': ref['cat2'],
            'category3': ref['cat3'],
            'brand': ref['brand'],
            'srp': ref['srp'],
            'f1': f1,
            'f2': f2,
            'total': f1 + f2,
            'stock_f1': f1,
            'stock_f2': f2,
            'stock_total': f1 + f2,
            'sourceLocation': 'BOTH_FLOORS' if (r1 and r2) else ('FLOOR_1_ONLY' if r1 else 'FLOOR_2_ONLY')
        })
        
    print(f"Total merged products: {len(merged_items)}")
    print(f"Grand F1: {f1_total} | Grand F2: {f2_total} | Total: {f1_total + f2_total}")
    
    colored_count = sum(1 for it in merged_items if it.get('color'))
    print(f"Products with resolved color: {colored_count} / {len(merged_items)}")
    
    lines = [
        '/**',
        ' * Samsung Branch Operations - Pilot Stock Snapshot',
        ' * Source: stock(1).xlsx (Sheet1: Floor 1 / f1, Sheet2: Floor 2 / f2)',
        f' * Total Products: {len(merged_items)} | F1: {f1_total} | F2: {f2_total} | Grand Total: {f1_total + f2_total}',
        ' */',
        'window.LATEST_STOCK_SNAPSHOT = ' + json.dumps(merged_items, ensure_ascii=False, indent=2) + ';',
        '',
        'window.STOCK_DATABASE = window.LATEST_STOCK_SNAPSHOT;',
        'window.STOCK_DATA = window.LATEST_STOCK_SNAPSHOT;',
        'window.STOCK_METADATA = {',
        '  stockBatchId: "STOCK-20260914-LATEST",',
        '  importBatchId: "STOCK-20260914-LATEST",',
        '  sourceType: "Manual Excel Snapshot",',
        '  sourceFilename: "stock(1).xlsx",',
        f'  recordCount: {len(merged_items)},',
        f'  uniquePn: {len(merged_items)},',
        f'  f1Total: {f1_total},',
        f'  f2Total: {f2_total},',
        f'  grandTotal: {f1_total + f2_total}',
        '};',
        ''
    ]
    
    with open(OUTPUT_JS, 'w', encoding='utf-8') as f:
        f.write('\n'.join(lines))
        
    print(f"Successfully generated {OUTPUT_JS}")

if __name__ == '__main__':
    main()
