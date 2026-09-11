# -*- coding: utf-8 -*-
"""
Samsung Branch Operations - Two-Sheet Excel Stock Parser
Parses Stock.xlsx with Sheet1 (Floor 1) and Sheet2 (Floor 2).
Executes strict Full Outer Join on Exact P/N.
Enforces:
- Quantity = On Hand only (On B/R, On Alloc, On T/F are metadata)
- total = f1 + f2
- Price 99 as stockReferencePrice (never overwrites promotion prices)
"""

import os
import sys
import json
import hashlib
import openpyxl
from datetime import datetime

sys.stdout.reconfigure(encoding='utf-8')

def compute_file_sha256(filepath):
    with open(filepath, "rb") as f:
        return hashlib.sha256(f.read()).hexdigest()

def parse_stock_excel(file_path):
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"Stock file not found: {file_path}")

    file_size = os.path.getsize(file_path)
    file_sha = compute_file_sha256(file_path)
    wb = openpyxl.load_workbook(file_path, data_only=True)

    required_sheets = ["Sheet1", "Sheet2"]
    for s in required_sheets:
        if s not in wb.sheetnames:
            raise ValueError(f"Missing required sheet '{s}' in workbook. Available: {wb.sheetnames}")

    sheet1 = wb["Sheet1"]
    sheet2 = wb["Sheet2"]

    def read_sheet_records(sheet, sheet_name):
        records = {}
        duplicates = []
        invalid_rows = []
        row_count = 0

        # Scan for header row containing 'P/N' and 'On Hand'
        header_row_idx = None
        col_map = {}
        for r_idx, row in enumerate(sheet.iter_rows(max_row=10, values_only=True), start=1):
            row_str = [str(c).strip() if c is not None else "" for c in row]
            if "P/N" in row_str and "On Hand" in row_str:
                header_row_idx = r_idx
                for c_idx, name in enumerate(row_str):
                    if name:
                        col_map[name] = c_idx
                break

        if not header_row_idx:
            raise ValueError(f"Cannot locate valid header row containing 'P/N' and 'On Hand' in {sheet_name}")

        pn_idx = col_map.get("P/N")
        on_hand_idx = col_map.get("On Hand")
        cat1_idx = col_map.get("Cat1")
        cat2_idx = col_map.get("Cat2")
        cat3_idx = col_map.get("Cat3")
        brand_idx = col_map.get("Brand")
        price99_idx = col_map.get("Price 99")
        koan_idx = col_map.get("KOAN SKU")
        apple_idx = col_map.get("Apple Part")
        desc_idx = col_map.get("Description")
        br_idx = col_map.get("On B/R")
        alloc_idx = col_map.get("On Alloc")
        tf_idx = col_map.get("On T/F")

        for r_idx, row in enumerate(sheet.iter_rows(min_row=header_row_idx + 1, values_only=True), start=header_row_idx + 1):
            if not any(row):
                continue
            row_count += 1
            raw_pn = row[pn_idx] if pn_idx < len(row) else None
            if not raw_pn or str(raw_pn).strip() == "":
                invalid_rows.append({"row": r_idx, "error": "MISSING_PN", "data": str(row[:5])})
                continue

            clean_pn = str(raw_pn).strip()
            raw_qty = row[on_hand_idx] if on_hand_idx < len(row) else None
            if raw_qty is None or str(raw_qty).strip() == "":
                invalid_rows.append({"row": r_idx, "pn": clean_pn, "error": "MISSING_ON_HAND"})
                continue

            try:
                qty = int(raw_qty)
                if qty < 0:
                    invalid_rows.append({"row": r_idx, "pn": clean_pn, "error": "NEGATIVE_ON_HAND", "value": qty})
                    continue
            except (ValueError, TypeError):
                invalid_rows.append({"row": r_idx, "pn": clean_pn, "error": "INVALID_NUMERIC_ON_HAND", "value": str(raw_qty)})
                continue

            if clean_pn in records:
                duplicates.append({"row": r_idx, "pn": clean_pn, "previousRow": records[clean_pn]["row"]})
                continue

            desc = str(row[desc_idx]).strip() if desc_idx is not None and desc_idx < len(row) and row[desc_idx] is not None else ""
            brand = str(row[brand_idx]).strip() if brand_idx is not None and brand_idx < len(row) and row[brand_idx] is not None else "UNKNOWN"
            cat1 = str(row[cat1_idx]).strip() if cat1_idx is not None and cat1_idx < len(row) and row[cat1_idx] is not None else ""
            cat2 = str(row[cat2_idx]).strip() if cat2_idx is not None and cat2_idx < len(row) and row[cat2_idx] is not None else ""
            cat3 = str(row[cat3_idx]).strip() if cat3_idx is not None and cat3_idx < len(row) and row[cat3_idx] is not None else ""
            price99 = row[price99_idx] if price99_idx is not None and price99_idx < len(row) else None

            records[clean_pn] = {
                "row": r_idx,
                "pn": clean_pn,
                "onHand": qty,
                "description": desc,
                "brand": brand,
                "category1": cat1,
                "category2": cat2,
                "category3": cat3,
                "stockReferencePrice": price99,
                "koanSku": str(row[koan_idx]).strip() if koan_idx is not None and koan_idx < len(row) and row[koan_idx] is not None else "",
                "applePart": str(row[apple_idx]).strip() if apple_idx is not None and apple_idx < len(row) and row[apple_idx] is not None else "",
                "onBackReserve": row[br_idx] if br_idx is not None and br_idx < len(row) else 0,
                "onAllocated": row[alloc_idx] if alloc_idx is not None and alloc_idx < len(row) else 0,
                "onTransfer": row[tf_idx] if tf_idx is not None and tf_idx < len(row) else 0
            }

        return records, duplicates, invalid_rows, row_count

    s1_records, s1_dups, s1_invalids, s1_rows = read_sheet_records(sheet1, "Sheet1 (Floor 1)")
    s2_records, s2_dups, s2_invalids, s2_rows = read_sheet_records(sheet2, "Sheet2 (Floor 2)")

    # Full Outer Join by Exact P/N
    all_pns = sorted(set(s1_records.keys()) | set(s2_records.keys()))
    matched_pns = set(s1_records.keys()) & set(s2_records.keys())
    s1_only_pns = set(s1_records.keys()) - set(s2_records.keys())
    s2_only_pns = set(s2_records.keys()) - set(s1_records.keys())

    joined_items = []
    total_f1 = 0
    total_f2 = 0
    total_stock = 0
    price_discrepancies = []

    for pn in all_pns:
        r1 = s1_records.get(pn)
        r2 = s2_records.get(pn)

        f1 = r1["onHand"] if r1 else 0
        f2 = r2["onHand"] if r2 else 0
        tot = f1 + f2

        total_f1 += f1
        total_f2 += f2
        total_stock += tot

        base = r1 if r1 else r2
        p1 = r1.get("stockReferencePrice") if r1 else None
        p2 = r2.get("stockReferencePrice") if r2 else None

        if p1 is not None and p2 is not None and p1 != p2:
            price_discrepancies.append({
                "pn": pn,
                "description": base.get("description"),
                "sheet1Price": p1,
                "sheet2Price": p2
            })

        # Determine inventory classification scope
        cat1 = base.get("category1", "")
        brand = base.get("brand", "")
        if "Smart Phone" in cat1:
            scope = "CORE_DEVICE"
        elif "Tablet" in cat1:
            scope = "CORE_DEVICE"
        elif brand.upper() == "SAMSUNG":
            scope = "SAMSUNG_ACCESSORY"
        elif "Accessory" in cat1 or "Audio" in cat1:
            scope = "THIRD_PARTY_ACCESSORY"
        elif "Service" in cat1:
            scope = "SIM_SERVICE"
        else:
            scope = "OTHER"

        joined_items.append({
            "pn": pn,
            "description": base.get("description"),
            "category1": cat1,
            "category2": base.get("category2"),
            "category3": base.get("category3"),
            "brand": brand,
            "stockReferencePrice": p1 if p1 is not None else p2,
            "f1": f1,
            "f2": f2,
            "total": tot,
            "inventoryScope": scope,
            "inSheet1": bool(r1),
            "inSheet2": bool(r2),
            "onBackReserve": (r1.get("onBackReserve", 0) if r1 else 0) + (r2.get("onBackReserve", 0) if r2 else 0),
            "onAllocated": (r1.get("onAllocated", 0) if r1 else 0) + (r2.get("onAllocated", 0) if r2 else 0),
            "onTransfer": (r1.get("onTransfer", 0) if r1 else 0) + (r2.get("onTransfer", 0) if r2 else 0)
        })

    batch_id = f"STOCK-BATCH-{datetime.now().strftime('%Y%m%d-%H%M%S')}"

    result = {
        "batchId": batch_id,
        "sourceFile": os.path.basename(file_path),
        "fileSizeBytes": file_size,
        "fileSha256": file_sha,
        "parsedAt": datetime.now().isoformat(),
        "summary": {
            "sheet1RowsRead": s1_rows,
            "sheet2RowsRead": s2_rows,
            "uniquePnSheet1": len(s1_records),
            "uniquePnSheet2": len(s2_records),
            "totalUniquePn": len(all_pns),
            "matchedPnBothSheets": len(matched_pns),
            "sheet1OnlyPn": len(s1_only_pns),
            "sheet2OnlyPn": len(s2_only_pns),
            "duplicateCountSheet1": len(s1_dups),
            "duplicateCountSheet2": len(s2_dups),
            "invalidRowCountSheet1": len(s1_invalids),
            "invalidRowCountSheet2": len(s2_invalids),
            "f1TotalStock": total_f1,
            "f2TotalStock": total_f2,
            "grandTotalStock": total_stock,
            "priceDiscrepancyCount": len(price_discrepancies)
        },
        "priceDiscrepancies": price_discrepancies,
        "items": joined_items
    }
    return result

if __name__ == "__main__":
    default_path = r"C:\Users\JarNJay\Desktop\Stock.xlsx"
    target = sys.argv[1] if len(sys.argv) > 1 else default_path
    print(f"Parsing: {target}")
    res = parse_stock_excel(target)
    print(f"✅ Success! Batch ID: {res['batchId']}")
    print(f"   Unique SKUs: {res['summary']['totalUniquePn']}")
    print(f"   F1 Total: {res['summary']['f1TotalStock']} | F2 Total: {res['summary']['f2TotalStock']} | Grand Total: {res['summary']['grandTotalStock']}")
