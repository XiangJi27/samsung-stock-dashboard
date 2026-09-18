import openpyxl
import sys
sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('Aug_ 2026 Promotion Retail_Shop Samsung .xlsx', data_only=True)
print('Sheet names:', wb.sheetnames)
ws = wb.active
for r in range(1, 35):
    row_vals = [str(c.value) if c.value is not None else '' for c in ws[r][:12]]
    if any(row_vals):
        print(f'Row {r}:', [v[:30] for v in row_vals])
