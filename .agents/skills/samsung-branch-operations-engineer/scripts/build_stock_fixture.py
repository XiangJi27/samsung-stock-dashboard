#!/usr/bin/env python3
"""
Automated Stock Snapshot Acceptance Fixture Builder
Generates hash-bound acceptance fixtures from the current Excel stock master.
Ensures stock figures are NEVER hardcoded as permanent business rules, but are
instead bound to the SHA-256 hash of the specific source file.
"""

import os
import sys
import json
import hashlib
import datetime
import openpyxl

sys.stdout.reconfigure(encoding="utf-8")

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
SKILL_DIR = os.path.dirname(SCRIPT_DIR)
WORKSPACE_ROOT = os.path.dirname(os.path.dirname(os.path.dirname(SKILL_DIR)))
DEFAULT_EXCEL = os.path.join(WORKSPACE_ROOT, "stock(1).xlsx")
OUTPUT_FIXTURE = os.path.join(SKILL_DIR, "fixtures", "stock_snapshot_acceptance.json")

def compute_sha256(filepath):
    hasher = hashlib.sha256()
    with open(filepath, "rb") as f:
        while chunk := f.read(65536):
            hasher.update(chunk)
    return hasher.hexdigest()

def classify_product(item):
    c1 = item['cat1']
    c2 = item['cat2']
    c3 = item['cat3']
    brand = item['brand']
    desc = item['desc']
    pn = item['pn']
    
    # 1. Galaxy Buds (MUST be evaluated before generic smartphone)
    if (c1 == 'AUDIO' and c2 == 'HEADPHONE' and c3 == 'TRUE WIRELESS' and 'SAMSUNG' in brand) or \
       ('SAMSUNG' in brand and (pn.startswith(('SM-R4', 'SM-R5', 'SM-R6')) or 'BUDS' in desc)):
        return 'Buds'
        
    # 2. Smartphone
    if c1 in ('SMART PHONES', 'SMARTPHONES', 'SMART PHONE', 'SMART_PHONES', 'SMART_PHONE'):
        return 'SmartPhone'
        
    # 3. Tablet
    if c1 in ('COMPUTER AND TABLET', 'COMPUTER_AND_TABLET', 'TABLET', 'TABLETS', 'TAB'):
        return 'Tablet'
        
    # 4. Smart Watch
    if c1 in ('SMART WATCH', 'SMART_WATCH', 'SMARTWATCH', 'WATCH'):
        return 'Watch'
        
    # 5. Accessories
    if c1 in ('MOBILE AND COMPUTER ACCESSORY', 'MOBILE_AND_COMPUTER_ACCESSORY', 'ACCESSORY', 'ACCESSORIES', 'ADAPTER'):
        return 'Accessory'
        
    # 6. Premium
    if 'PREMIUM' in c1 or 'PREMIUM' in c2 or 'FREE GIFT' in c2 or 'PREMIUM' in desc or 'FREE GIFT' in desc or 'GIFT' in c1:
        return 'Premium'
        
    # 7. SIM
    if 'SERVICE, INSURANCE AND WARRANTY' in c1 or 'SERVICE,_INSURANCE_AND_WARRANTY' in c1 or 'SIM' in c1 or 'SIM' in c2 or 'CARRIER MOBILE PACKAGE' in c2:
        return 'SIM'
        
    return 'Other'

def parse_sheet(ws):
    rows = []
    for r in range(4, ws.max_row + 1):
        cat1 = str(ws.cell(row=r, column=1).value or '').strip().upper()
        cat2 = str(ws.cell(row=r, column=2).value or '').strip().upper()
        cat3 = str(ws.cell(row=r, column=3).value or '').strip().upper()
        brand = str(ws.cell(row=r, column=4).value or '').strip().upper()
        pn = str(ws.cell(row=r, column=6).value or '').strip().upper()
        desc = str(ws.cell(row=r, column=9).value or '').strip().upper()
        on_hand = ws.cell(row=r, column=10).value
        
        if not pn or pn == 'TOTAL' or cat1 == 'TOTAL':
            continue
            
        try:
            qty = int(on_hand or 0)
        except Exception:
            qty = 0
            
        rows.append({
            'pn': pn,
            'cat1': cat1,
            'cat2': cat2,
            'cat3': cat3,
            'brand': brand,
            'desc': desc,
            'qty': qty
        })
    return rows

def build_fixture(excel_path=DEFAULT_EXCEL, output_path=OUTPUT_FIXTURE):
    if not os.path.exists(excel_path):
        raise FileNotFoundError(f"Excel file not found at: {excel_path}")
        
    file_sha256 = compute_sha256(excel_path)
    filename = os.path.basename(excel_path)
    
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws1 = wb['Sheet1']
    ws2 = wb['Sheet2']
    
    s1_rows = parse_sheet(ws1)
    s2_rows = parse_sheet(ws2)
    
    s1_map = {r['pn']: r for r in s1_rows}
    s2_map = {r['pn']: r for r in s2_rows}
    all_pns = sorted(list(set(list(s1_map.keys()) + list(s2_map.keys()))))
    
    f1_cat_counts = {}
    f1_cat_pns = {}
    f1_total = 0
    f2_total = 0
    
    for r in s1_rows:
        cat = classify_product(r)
        f1_cat_counts[cat] = f1_cat_counts.get(cat, 0) + r['qty']
        f1_cat_pns.setdefault(cat, set()).add(r['pn'])
        f1_total += r['qty']
        
    for r in s2_rows:
        f2_total += r['qty']
        
    fixture_data = {
        "fixtureType": "SNAPSHOT_ACCEPTANCE",
        "sourceFilename": filename,
        "sourceSha256": file_sha256,
        "generatedAt": datetime.datetime.now(datetime.timezone(datetime.timedelta(hours=7))).isoformat(),
        "expected": {
            "f1Total": f1_total,
            "smartphone": f1_cat_counts.get('SmartPhone', 0),
            "tablet": f1_cat_counts.get('Tablet', 0),
            "watch": f1_cat_counts.get('Watch', 0),
            "buds": f1_cat_counts.get('Buds', 0),
            "accessory": f1_cat_counts.get('Accessory', 0),
            "premium": f1_cat_counts.get('Premium', 0)
        },
        "inventoryBreakdown": {
            "sim": {
                "quantityUnits": f1_cat_counts.get('SIM', 0),
                "distinctPnItems": len(f1_cat_pns.get('SIM', set())),
                "unitLabel": "ชิ้น",
                "pnLabel": "รายการ"
            },
            "other": {
                "quantityUnits": f1_cat_counts.get('Other', 0),
                "distinctPnItems": len(f1_cat_pns.get('Other', set())),
                "unitLabel": "ชิ้น",
                "pnLabel": "รายการ"
            },
            "f2Total": f2_total,
            "grandTotal": f1_total + f2_total,
            "totalProducts": len(all_pns)
        },
        "rulesEnforced": {
            "summaryCardsScope": "F1_ONLY",
            "suppressedCards": ["SIM", "OTHER"],
            "mathematicalCheck": f"{f1_cat_counts.get('SmartPhone', 0) + f1_cat_counts.get('Tablet', 0) + f1_cat_counts.get('Watch', 0) + f1_cat_counts.get('Buds', 0) + f1_cat_counts.get('Accessory', 0) + f1_cat_counts.get('Premium', 0)} (Hardware + Premium) + {f1_cat_counts.get('SIM', 0)} (SIM) + {f1_cat_counts.get('Other', 0)} (Other) == {f1_total} (Total F1)"
        }
    }
    
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(fixture_data, f, ensure_ascii=False, indent=2)
        
    print(f"✅ Generated Snapshot Acceptance Fixture:")
    print(f"   Source File: {filename}")
    print(f"   SHA-256:     {file_sha256}")
    print(f"   F1 Total:    {f1_total}")
    print(f"   SIM:         {f1_cat_counts.get('SIM', 0)} ชิ้น ({len(f1_cat_pns.get('SIM', set()))} รายการ P/N)")
    print(f"   Other:       {f1_cat_counts.get('Other', 0)} ชิ้น ({len(f1_cat_pns.get('Other', set()))} รายการ P/N)")
    print(f"   Output:      {output_path}")
    return fixture_data

if __name__ == "__main__":
    target_excel = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_EXCEL
    build_fixture(target_excel)
